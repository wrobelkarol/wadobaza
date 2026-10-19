#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Podpowiedź Excela: z pliku w starym (lub częściowo przerobionym) układzie robi NOWY plik w docelowym układzie.

Oryginał zostaje nietknięty. Nowy plik powstaje w tym samym folderze, z dopiskiem _nowy
(1890.xlsx -> 1890_nowy.xlsx; gdy taki już jest: _nowy2, _nowy3 ...).

Co robi z każdym aktem:
  * opis aktu: bez łacińskich nazw w nawiasach, miejscowości w miejscowniku (ur. w Przytkowicach)
  * Imię nazwisko, Mama, Tata, Małżonek: z Twojego pliku, jeśli już wpisane; w wierszach jeszcze
    nieprzerobionych – wstępnie z opisu aktu (komórki zaznaczone na żółto = do sprawdzenia)
  * kobiety pod nazwiskiem panieńskim; gdy opis go nie podaje (ani wprost, ani przez ojca) – bez nazwiska
  * Uwagi (bez kapłana) i Kapłan (kto zaopatrzył i pochował) – rozdzielone
  * wyznanie, płeć, wiek, przyczyna zgonu, data urodzenia, skan – przepisane
Żółtą komórkę po sprawdzeniu popraw lub wyczyść wypełnienie. Zielony znacznik "sprawdzone" zostaje Twój.

Użycie:
  python skrypty/podpowiedz_excel.py "C:/.../Przytkowice/Zgony/1890/1890.xlsx"
  python skrypty/podpowiedz_excel.py "C:/.../Przytkowice/Zgony/1890" "C:/.../Przytkowice/Zgony/1891"
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import math
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # żeby działało też z opcją -I
import build  # noqa: E402

try:
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import PatternFill
except ImportError:  # pragma: no cover
    sys.exit("Brak biblioteki openpyxl. Zainstaluj ją poleceniem: pip install openpyxl")

WZOR = Path(__file__).resolve().parent / "wzor_zgony.xlsx"
ZOLTY = PatternFill("solid", fgColor="FFF2CC")
BEZ_WYPELNIENIA = PatternFill(fill_type=None)


def format_daty(iso: str | None) -> str | None:
    """'1881-12-01' -> '1.12.1881'; '1941-04' -> '4.1941'; '1928' -> '1928'."""
    if not iso:
        return None
    if len(iso) == 10:
        return f"{int(iso[8:])}.{int(iso[5:7])}.{iso[:4]}"
    if len(iso) == 7:
        return f"{int(iso[5:])}.{iso[:4]}"
    return iso


def surowa_data(v: object) -> str | None:
    """Dzień zgonu/pogrzebu tak, jak był w Excelu (tekst zostaje, daty zamieniamy na d.m.rrrr)."""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.day}.{v.month}.{v.year}"
    return build.tekst(v)


def pelne(imie: str | None, nazwisko: str | None) -> str | None:
    return " ".join(x for x in (imie, nazwisko) if x) or None


def nowa_sciezka(zrodlo: Path) -> Path:
    cel = zrodlo.with_name(f"{zrodlo.stem}_nowy.xlsx")
    n = 2
    while cel.exists():
        cel = zrodlo.with_name(f"{zrodlo.stem}_nowy{n}.xlsx")
        n += 1
    return cel


def kolumny_wzoru(ws) -> dict:
    kol = build.wykryj_kolumny(ws)
    for c in kol["wiekowe"]:
        naglowek = build.zloz(ws.cell(1, c).value or "")
        kol["data_ur" if "urodzeni" in naglowek else "wiek"] = c
    for brak in ("osoba", "matka", "ojciec", "malzonek", "kaplan", "skan", "uwagi_kol", "wiek", "data_ur"):
        if brak != "uwagi_kol" and brak not in kol:
            sys.exit(f"Wzór {WZOR.name} nie ma kolumny: {brak}")
    kol["uwagi_kol"] = kol["uwagi"][0]
    return kol


def przygotuj_wiersz(rek: dict) -> dict:
    """Wartości dla kolumn nowego arkusza + które z nich są podpowiedzią (żółte)."""
    s = rek["_surowe"]
    jawny = not rek.get("rodzice_auto") and bool(s["jawna"] or rek.get("matka") or rek.get("ojciec") or rek.get("malzonek"))
    zolte: set[str] = set()
    wart: dict[str, object] = {}

    wart["nr"] = rek.get("nr")
    wart["data_zgonu"] = surowa_data(s["zgon"])
    wart["data_pogrzebu"] = surowa_data(s["pogrzeb"])
    wart["dom"] = rek.get("dom")
    wart["opis"] = rek.get("opis")

    if s["jawna"]:
        wart["osoba"] = s["jawna"]
    else:
        wart["osoba"] = pelne(rek.get("imie"), rek.get("nazwisko"))
        zolte.add("osoba")
    if jawny:  # wiersz już przerobiony ręcznie – przepisujemy tylko to, co wpisano
        wart["matka"], wart["ojciec"], wart["malzonek"] = rek.get("matka"), rek.get("ojciec"), rek.get("malzonek")
    else:
        wart["matka"] = pelne(rek.get("imie_matki"), rek.get("nazwisko_matki"))
        wart["ojciec"] = pelne(rek["imie_ojca"], rek.get("nazwisko")) if rek.get("imie_ojca") else None
        wart["malzonek"] = rek.get("malzonek")
        zolte.update(k for k in ("matka", "ojciec", "malzonek") if wart[k])
    wart["rel_kat"] = 1 if rek.get("religia") == "kat" else None
    wart["rel_inna"] = 1 if rek.get("religia") == "inna" else None
    wart["plec_m"] = 1 if rek.get("plec") == "M" else None
    wart["plec_k"] = 1 if rek.get("plec") == "K" else None
    wart["wiek"] = s["wiek_tekst"]
    wart["przyczyna"] = rek.get("przyczyna")
    wart["data_ur"] = format_daty(rek.get("data_ur"))
    wart["data_slubu"] = format_daty(rek.get("data_slubu"))  # wpisze się tylko, jeśli wzór ma kolumnę "Data ślubu"
    wart["uwagi_kol"] = rek.get("dodatkowe")
    wart["skan"] = rek.get("skan")
    wart["kaplan"] = rek.get("kaplan")
    if not (s["jawna"] or s["kaplan_jawny"]):  # wiersze nieprzerobione: uwagi i kapłan też są podpowiedzią
        zolte.update(k for k in ("uwagi_kol", "kaplan") if wart[k])
    return {"wart": wart, "zolte": zolte, "jawny": jawny}


def zrob_plik(zrodlo: Path, znane: set[str]) -> None:
    args = argparse.Namespace(max_rok=9999, bez_uwag=False)
    raport = build.Raport()
    rekordy = build.wczytaj_plik(zrodlo, args, raport, znane)
    if not rekordy:
        print(f"POMINIĘTO {zrodlo.name}: " + ("; ".join(raport.pliki_pominiete) or "brak aktów"))
        return

    wb = openpyxl.load_workbook(WZOR)
    ws = wb.worksheets[0]
    kol = kolumny_wzoru(ws)
    ostatnia = max(v for k, v in kol.items() if isinstance(v, int))
    rok = rekordy[0]["rok"]
    ws["B1"] = f"{rok}\nDzień i miesiąc"

    wzor_stylu = {c: copy.copy(ws.cell(3, c)._style) for c in range(1, ws.max_column + 1)}
    for r in range(3, 3 + len(rekordy)):
        if r > 3:
            ws.merge_cells(start_row=r, start_column=kol["opis"], end_row=r, end_column=kol["opis"] + 2)
        for c, stl in wzor_stylu.items():
            cell = ws.cell(r, c)
            cell._style = copy.copy(stl)
            if c > ostatnia:  # kolumny za ostatnią nazwaną (znacznik "sprawdzone") zostają puste i bez koloru
                cell.fill = BEZ_WYPELNIENIA
    # wiersz 3 we wzorze ma kolorowy znacznik za ostatnią kolumną – czyścimy go także tam
    for c in range(ostatnia + 1, ws.max_column + 1):
        ws.cell(3, c).fill = BEZ_WYPELNIENIA

    zolte_razem = 0
    for i, rek in enumerate(rekordy):
        r = 3 + i
        wiersz = przygotuj_wiersz(rek)
        for nazwa, v in wiersz["wart"].items():
            if v in (None, "") or nazwa not in kol:
                continue
            cell = ws.cell(r, kol[nazwa])
            if isinstance(cell, MergedCell):
                continue
            cell.value = v
            if nazwa in wiersz["zolte"]:
                cell.fill = ZOLTY
                zolte_razem += 1
        linie = math.ceil(len(rek.get("opis") or "") / 26) + 1
        ws.row_dimensions[r].height = max(60, min(409, linie * 14.5))

    cel = nowa_sciezka(zrodlo)
    wb.save(cel)
    print(f"OK  {zrodlo.name} -> {cel.name}  (aktów: {len(rekordy)}, żółtych podpowiedzi: {zolte_razem})")
    print(f"    {cel}")
    for o in raport.ostrzezenia[:5]:
        print("    !", o)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sciezki", nargs="+", type=Path, help="plik .xlsx albo folder z plikami")
    args = ap.parse_args()
    if not WZOR.exists():
        sys.exit(f"Brak wzoru arkusza: {WZOR}")
    pliki = build.zbierz_pliki(args.sciezki)
    if not pliki:
        sys.exit("Nie znaleziono żadnych plików .xlsx.")
    znane = build.zbierz_znane(pliki)
    for p in pliki:
        zrob_plik(p, znane)


if __name__ == "__main__":
    main()
