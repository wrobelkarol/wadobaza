#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Wadobaza: Excel (.xlsx) -> JSON dla strony.

Skrypt czyta przetłumaczone akty z plików Excel, sprawdza je, ujednolica
i zapisuje do docs/data/. Pliki Excel zostają u Ciebie na dysku, do repozytorium
trafiają tylko wygenerowane pliki JSON.

Użycie:
  python skrypty/build.py                          # źródła z pliku skrypty/zrodla.txt
  python skrypty/build.py --zrodlo "C:/.../Przytkowice/Zgony/1941/1941_2.xlsx"
  python skrypty/build.py --zrodlo "C:/.../Przytkowice/Zgony" --sucho --szczegoly

Struktura folderów źródłowych: <Parafia>/<Zgony|Urodzenia|Śluby>/<rok>/plik.xlsx
(parafię, rodzaj aktu i rok skrypt bierze ze ścieżki i nagłówka arkusza).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # żeby działało też z opcją -I
from opis import (  # noqa: E402
    dodatkowe_z_opisu, kobieta_zamezna, nazwisko_mianownik, nazwisko_panienskie, popraw_miejscowosci,
    rodzina_z_opisu, rozdziel_kaplana,
    usun_lacine, wiek_pelne_lata, zbierz_znane_nazwiska,
)

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Brak biblioteki openpyxl. Zainstaluj ją poleceniem: pip install openpyxl")

KORZEN = Path(__file__).resolve().parent.parent
KATALOG_WYJSCIA = KORZEN / "docs" / "data"
PLIK_ZRODEL_SKANOW = KORZEN / "dane" / "zrodla.json"
PLIK_LISTY_ZRODEL = Path(__file__).resolve().parent / "zrodla.txt"

MAX_ROK_DOMYSLNIE = 1945
MAX_ROK_TYPU = {"urodzenia": 1926}  # narodziny: osoby mogą jeszcze żyć, więc ostrzejszy limit niż dla pozostałych aktów
LICENCJA = {
    "nazwa": "CC BY-NC 4.0",
    "url": "https://creativecommons.org/licenses/by-nc/4.0/deed.pl",
}

# rodzaj aktu w nazwie folderu -> identyfikator
TYPY_FOLDEROW = {
    "zgony": "zgony",
    "zgon": "zgony",
    "urodzenia": "urodzenia",
    "chrzty": "urodzenia",
    "sluby": "malzenstwa",
    "malzenstwa": "malzenstwa",
}
TYPY_OBSLUGIWANE = {"zgony"}  # pozostałe dojdą, gdy będzie gotowy układ Excela

PUSTE = {"", "[puste]", "-", "—", "–", "none", "nan"}


# ---------------------------------------------------------------- narzędzia

def zloz(s: object) -> str:
    """Małe litery, bez polskich znaków i nadmiarowych spacji."""
    s = str(s).replace("ł", "l").replace("Ł", "L")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s).strip().lower()


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", zloz(s)).strip("-")


def tekst(v: object) -> str | None:
    """Wartość komórki jako czysty tekst; puste znaczniki -> None."""
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\r", "").strip()
    s = re.sub(r"[ \t\u00a0]+", " ", s)
    s = re.sub(r"\s*\n\s*", " ", s)
    return None if zloz(s) in PUSTE else s


def zolta(komorka) -> bool:
    """Czy komórka ma żółte wypełnienie z podpowiedz_excel.py (FFF2CC) – wartość do sprawdzenia."""
    f = komorka.fill
    return bool(f is not None and f.fill_type == "solid" and str(f.fgColor.rgb).upper().endswith("FFF2CC"))


def poprawione(t: str | None) -> str | None:
    """Tekst bez łacińskich nawiasów i z miejscowościami w miejscowniku."""
    return popraw_miejscowosci(usun_lacine(t)) if t else t


def liczba(v: object) -> int | None:
    t = tekst(v)
    if t is None:
        return None
    try:
        return int(float(t.replace(",", ".")))
    except ValueError:
        return None


def zaznaczone(v: object) -> bool:
    t = tekst(v)
    return t is not None and zloz(t) in {"1", "x", "tak", "t"}


# ---------------------------------------------------------------- daty i wiek

RE_ISO = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
RE_DATA = re.compile(r"(\d{1,2})\s*[.\-/]\s*(\d{1,2})\s*[.\-/]\s*(\d{4})")


def _data_lub_none(r: int, m: int, d: int) -> str | None:
    try:
        return dt.date(r, m, d).isoformat()
    except ValueError:
        return None


def parsuj_date(v: object) -> tuple[str | None, str | None]:
    """-> (data, problem). Data to 'RRRR-MM-DD' albo samo 'RRRR' (gdy znany tylko rok)."""
    if v is None:
        return None, None
    if isinstance(v, dt.datetime):
        v = v.date()
    if isinstance(v, dt.date):
        return v.isoformat(), None
    if isinstance(v, float) and not v.is_integer():  # "4.1941" wpisane jako liczba = kwiecień 1941 (bez dnia)
        m = re.fullmatch(r"(\d{1,2})\.(\d{4})", repr(v))
        if m and 1 <= int(m[1]) <= 12:
            return f"{m[2]}-{int(m[1]):02d}", None
    if isinstance(v, (int, float)):
        n = int(v)
        if 1700 <= n <= 2100:
            return str(n), None
        return None, f"liczba {v} zamiast daty"
    s = str(v).strip()
    f = zloz(s)
    if f in PUSTE or "nie widoczn" in f or f.startswith("brak"):
        return None, None
    m = RE_ISO.search(s)
    if m:
        wynik = _data_lub_none(int(m[1]), int(m[2]), int(m[3]))
        return (wynik, None) if wynik else (None, f"nieprawidłowa data '{s}'")
    m = RE_DATA.search(s)
    if m:
        wynik = _data_lub_none(int(m[3]), int(m[2]), int(m[1]))
        return (wynik, None) if wynik else (None, f"nieprawidłowa data '{s}'")
    mm = re.fullmatch(r"(\d{1,2})\.(\d{4})", s)  # miesiąc i rok bez dnia
    if mm and 1 <= int(mm[1]) <= 12:
        return f"{mm[2]}-{int(mm[1]):02d}", None
    if re.fullmatch(r"\d{4}", s):
        return s, None
    return None, f"nie rozpoznano daty '{s}'"


def rok_z_daty(d: str | None) -> int | None:
    return int(d[:4]) if d else None


def rozdziel_wiek(v: object) -> tuple[str | None, str | None]:
    """Komórka z kolumny wieku/urodzenia -> (data_urodzenia, tekst_wieku)."""
    if v is None:
        return None, None
    if isinstance(v, (dt.datetime, dt.date, int, float)):
        d, _ = parsuj_date(v)
        return d, None
    t = tekst(v)
    if t is None:
        return None, None
    if re.fullmatch(r"(ur\.?\s*)?(\d{1,2}\.\d{1,2}\.\d{4}|\d{4}-\d{2}-\d{2}.*|\d{4})", t.strip()):
        d, _ = parsuj_date(re.sub(r"^ur\.?\s*", "", t.strip()))
        return d, None
    ur = re.search(r"\bur\.?\s*(\d{1,2}\.\d{1,2}\.\d{4})", t)
    d = parsuj_date(ur[1])[0] if ur else None
    return d, t


def wiek_w_latach(t: str | None) -> int | None:
    if not t:
        return None
    m = re.match(r"\s*(\d+)(?:[.,]\d+)?\s*(?:lat|lata|rok|roku|l\.?)\b", t, re.I)
    if m:
        return int(m[1])
    if re.match(r"\s*\d+\s*(?:miesi|mies|dni|dzie|tydz|tyg|godz)", t, re.I):
        return 0
    return None


# ---------------------------------------------------------------- opis aktu

CAP = "A-ZĄĆĘŁŃÓŚŹŻ"
LOW = "a-ząćęłńóśźż"
SLOWO = rf"[{CAP}][{LOW}{CAP}'’\-]*"

RE_IMIE_NAZWISKO = re.compile(
    rf"^\s*(?P<imie>{SLOWO}(?:\s+{SLOWO})*)\s*\((?P<lac>[^)]*)\)"
    rf"(?:\s*\([^)]*\))*"
    rf"\s*(?P<nazw>{SLOWO})?"
)
RE_RODZIC = re.compile(
    rf"\b(?:syn|córka|córką|synem|dziecko)\s+(?P<niesl>nieślubn\w+\s+)?"
    rf"(?P<oj>{SLOWO}(?:\s+{SLOWO})*)\s*\((?P<lac>[^)]*)\)(?:\s+(?P<nazw>{SLOWO}))?"
    rf"(?:\s+z\s+domu\s+(?P<pan>{SLOWO}))?"
)
RE_BEZ_LACINY = re.compile(rf"^\s*(?P<imie>{SLOWO}(?:\s+{SLOWO})?)\s+(?P<nazw>{SLOWO})\s*(?:,|\.|\[|$)")
RE_RODZIC_BEZ_LACINY = re.compile(
    rf"\b(?:syn|córka)\s+(?P<niesl>nieślubn\w+\s+)?(?P<oj>{SLOWO})\s+(?P<nazw>{SLOWO})(?:\s*\((?P<nom>[^)]*)\))?"
)
RE_PO_NAWIASIE = re.compile(rf"\)\s*(?P<oj>{SLOWO})\s*\([^)]*\)\s*(?P<nazw>{SLOWO})")
RE_MARTWO = re.compile(r"martwo|nieżywo|nieochrzczon", re.I)
RE_MIEJSCE = re.compile(r"\bur\.\s+(?:w|we|na)\s+(?P<m>[^.,(\[;]+)")


def nazwisko_od_rodzica(opis: str) -> str | None:
    """Dziecko bez nazwiska w nagłówku aktu: bierzemy nazwisko ojca (lub panieńskie matki)."""
    r = RE_RODZIC.search(opis)
    if r:
        if r["nazw"]:
            return r["nazw"]
        if (r["niesl"] or re.search(r"nieślubn\w+\s+(?:syn|córka)\b", opis)) and r["pan"]:
            return r["pan"]
    r = RE_RODZIC_BEZ_LACINY.search(opis)
    if r:
        nom = (r["nom"] or "").strip()
        return nom if re.fullmatch(SLOWO, nom) else r["nazw"]
    if RE_MARTWO.search(opis[:80]):
        r = RE_PO_NAWIASIE.search(opis)
        if r:
            return r["nazw"]
    return None


def rozbierz_opis(opis: str) -> dict:
    """Wyciąga z opisu zmarłego: imię, nazwisko, miejsce urodzenia (heurystyka)."""
    wynik: dict = {}
    m = RE_IMIE_NAZWISKO.match(opis) or RE_BEZ_LACINY.match(opis)
    if m:
        wynik["imie"] = m["imie"].strip()
        if m["nazw"]:
            wynik["nazwisko"] = m["nazw"].strip()
    if "nazwisko" not in wynik:
        nazw = nazwisko_od_rodzica(opis)
        if nazw:
            wynik["nazwisko"] = nazw
            wynik["nazwisko_wg_rodzica"] = True
    if "imie" not in wynik:
        glowa = re.split(r"[,(\[]", opis, maxsplit=1)[0].strip()
        if glowa and len(glowa) < 50 and re.match(rf"[{CAP}]", glowa):
            wynik["imie"] = glowa
    mm = RE_MIEJSCE.search(opis)
    if mm:
        miejsce = mm["m"].strip()
        if miejsce:
            wynik["miejsce_ur"] = miejsce
    return wynik


def rozdziel_imie_nazwisko(s: str) -> tuple[str | None, str | None]:
    """'Marianna Cichoń' -> ('Marianna', 'Cichoń'); 'Jan Kanty Nowak' -> ('Jan Kanty', 'Nowak')."""
    slowa = s.split()
    if len(slowa) == 1:
        return slowa[0], None
    return " ".join(slowa[:-1]), slowa[-1]


def normalizuj_skan(v: object) -> str | None:
    """Numer zdjęcia w Skanotece: 2 / '002' / '002.JPG' -> '002.JPG'."""
    t = tekst(v)
    if not t:
        return None
    m = re.fullmatch(r"(\d{1,4})(?:\.(?:jpe?g|png))?", t.strip(), re.I)
    return f"{int(m[1]):03d}.JPG" if m else t.strip()


# ---------------------------------------------------------------- arkusz Excel

class BladUkladu(Exception):
    pass


def wykryj_kolumny(ws) -> dict:
    """Rozpoznaje kolumny po nagłówkach z wierszy 1-2 (z uwzględnieniem scalonych komórek)."""
    scalone: dict[tuple[int, int], object] = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= 2:
            wartosc = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, min(rng.max_row, 2) + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    scalone[(r, c)] = wartosc

    def naglowek(r: int, c: int) -> str:
        v = scalone.get((r, c), ws.cell(r, c).value)
        return zloz(v) if v is not None else ""

    kol: dict[str, int] = {}
    wiekowe: list[int] = []
    uwagi_jawne: list[int] = []
    ostatnia_nazwana = 0
    for c in range(1, ws.max_column + 1):
        h1, h2 = naglowek(1, c), naglowek(2, c)
        wiersz_naglowka = f"{h1} {h2}"
        nazwa = None
        if "liczba porz" in wiersz_naglowka:
            nazwa = "nr"
        elif "numer domu" in wiersz_naglowka:
            nazwa = "dom"
        elif h1.startswith("imie, nazwisko"):
            nazwa = "opis"
        elif h1 in ("imie nazwisko", "imie i nazwisko", "zmarly", "zmarla"):
            nazwa = "osoba"
        elif h1 in ("mama", "matka"):
            nazwa = "matka"
        elif h1 in ("tata", "ojciec"):
            nazwa = "ojciec"
        elif h1.startswith("malzon"):
            nazwa = "malzonek"
        elif "skan" in h1:
            nazwa = "skan"
        elif h1.startswith("kaplan"):
            nazwa = "kaplan"
        elif "slub" in h1:  # "Data ślubu" – adnotacja dopisana później w księdze zgonów
            nazwa = "data_slubu"
        elif h2 == "zgonu":
            nazwa = "data_zgonu"
        elif "pogrzeb" in h2:
            nazwa = "data_pogrzebu"
        elif h2 == "katolickie":
            nazwa = "rel_kat"
        elif h2 == "inne":
            nazwa = "rel_inna"
        elif h2 == "meska":
            nazwa = "plec_m"
        elif h2 == "zenska":
            nazwa = "plec_k"
        elif "przyczyna" in h1:
            nazwa = "przyczyna"
        elif "wiek" in h1 or "urodzeni" in h1 or "narodzin" in h1:
            wiekowe.append(c)
            ostatnia_nazwana = max(ostatnia_nazwana, c)
            continue
        elif any(s in h1 for s in ("uwag", "komentarz", "notatk")):
            uwagi_jawne.append(c)
            ostatnia_nazwana = max(ostatnia_nazwana, c)
            continue
        if nazwa and nazwa not in kol:
            kol[nazwa] = c
            ostatnia_nazwana = max(ostatnia_nazwana, c)
    kol["wiekowe"] = wiekowe  # type: ignore[assignment]
    # Uwagi: kolumna z nagłówkiem "Uwagi"; w starszym układzie – kolumny bez nagłówka za ostatnią nazwaną.
    # Kolumna bez nagłówka obok jawnych "Uwag" (np. zielony znacznik "sprawdzone") jest ignorowana.
    kol["uwagi"] = uwagi_jawne or [  # type: ignore[assignment]
        c for c in range(ostatnia_nazwana + 1, ws.max_column + 1) if not naglowek(1, c)
    ]
    brak = [n for n in ("nr", "opis") if n not in kol]
    if "data_zgonu" not in kol and "data_pogrzebu" not in kol:
        brak.append("data zgonu / pogrzebu")
    if brak:
        raise BladUkladu("nie rozpoznano kolumn: " + ", ".join(brak))
    return kol


def rozpoznaj_sciezke(plik: Path) -> tuple[str | None, str | None, int | None]:
    """Parafia, rodzaj aktu i rok na podstawie ścieżki <Parafia>/<Rodzaj>/<rok>/plik."""
    czesci = plik.parts
    for i, czesc in enumerate(czesci[:-1]):
        typ = TYPY_FOLDEROW.get(zloz(czesc))
        if typ and i > 0:
            rok = None
            if i + 1 < len(czesci) - 1 and re.fullmatch(r"\d{4}", czesci[i + 1]):
                rok = int(czesci[i + 1])
            return czesci[i - 1], typ, rok
    return None, None, None


class Raport:
    def __init__(self) -> None:
        self.ostrzezenia: list[str] = []
        self.bez_nazwiska: list[str] = []
        self.odrzucone_rok: Counter = Counter()
        self.pliki_pominiete: list[str] = []

    def ostrz(self, plik: str, wiersz: int | None, komunikat: str) -> None:
        gdzie = f"{plik}" + (f" w.{wiersz}" if wiersz else "")
        self.ostrzezenia.append(f"{gdzie}: {komunikat}")


def zbierz_znane(pliki: list[Path]) -> set[str]:
    """Nazwiska w mianowniku znane z danych (potrzebne do odmiany nazwisk z opisów: Kukuły -> Kukuła)."""
    teksty: list[str] = []
    znane: set[str] = set()
    for p in pliki:
        if rozpoznaj_sciezke(p)[1] not in TYPY_OBSLUGIWANE:
            continue
        try:
            ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
            kol = wykryj_kolumny(ws)
        except Exception:  # plik z innym układem zgłosi błąd przy właściwym wczytaniu
            continue
        for w in range(3, ws.max_row + 1):
            o = tekst(ws.cell(w, kol["opis"]).value)
            if o:
                teksty.append(o)
            for nazwa in ("osoba", "matka", "ojciec", "malzonek"):
                v = tekst(ws.cell(w, kol[nazwa]).value) if nazwa in kol else None
                if v and len(v.split()) >= 2:
                    znane.add(v.split()[-1])
    return znane | zbierz_znane_nazwiska(teksty)


def czy_weryfikacja(wb) -> bool:
    """Plik pobrany z programu Weryfikacja: arkusze 'Łacina' i 'Polski', w każdym kolumna 'Numer aktu'."""
    return "Polski" in wb.sheetnames and "Numer aktu" in [c.value for c in wb["Polski"][1]]


def _dzien(v: object) -> tuple[int | None, int | None]:
    """Dzień z Weryfikacji: 5 / '5.' -> (5, None); '28/2' -> (28, 2)."""
    m = re.fullmatch(r"\s*(\d{1,2})\s*(?:/\s*(\d{1,2}))?\s*\.?\s*", str(v if v is not None else ""))
    return (int(m[1]), int(m[2]) if m[2] else None) if m else (None, None)


def _bez_ulamka_lat(t: str | None) -> str | None:
    """'40½ lat' / '4 9/12 lat' -> '40 lat' / '4 lat' (strona pokazuje pełne lata)."""
    return re.sub(r"^\s*(\d+)(?:\s*[½¼¾⅓⅔]|\s+\d+\s*/\s*\d+)(?=\s*(?:lat|lata|rok|roku)\b)", r"\1", t) if t else t


def forma_zenska(nazwisko: str) -> str:
    """Nazwisko ojca dla córki: Rudecki -> Rudecka, Ciepły -> Ciepła; rzeczownikowe (Rzepka, Morek) bez zmian."""
    if re.search(r"(sk|ck|dzk)i$", nazwisko) or re.search(r"[^aeiouyąęó]y$", nazwisko):
        return nazwisko[:-1] + "a"
    return nazwisko


def wczytaj_weryfikacje(wb, plik: Path, parafia: str, typ: str, rok_folder: int | None, args, raport: Raport) -> list[dict]:
    """Rocznik z programu Weryfikacja: każde pole aktu w osobnej kolumnie, sprawdzone ręcznie (bez zgadywania z opisu)."""
    ws = wb["Polski"]
    naglowki = [tekst(c.value) for c in ws[1]]
    nazwa_pliku = f"{parafia}/{plik.parent.name}/{plik.name}"

    rekordy: list[dict] = []
    for w, wiersz in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        k = {h: wiersz[i] for i, h in enumerate(naglowki) if h}
        t = lambda kol: tekst(k.get(kol))  # noqa: E731
        nr, opis = liczba(k.get("Numer aktu")), t("Tekst aktu (cały)")
        if nr is None and opis is None:
            continue
        if opis is None:
            raport.ostrz(nazwa_pliku, w, f"akt nr {nr} bez tekstu aktu – pominięty")
            continue
        rok = liczba(k.get("Rok")) or rok_folder
        if not rok:
            raport.ostrz(nazwa_pliku, w, "nie ustalono roku – pominięty")
            continue

        # daty: dzień + miesiąc (miesiąc pogrzebu: kolumna, zapis "3/2" albo następny, gdy dzień pogrzebu jest wcześniejszy)
        dz, mz = _dzien(k.get("Dzień zgonu"))
        dp, mp = _dzien(k.get("Dzień pogrzebu"))
        m_zgonu = mz or liczba(k.get("Miesiąc zgonu"))
        m_pogrz = liczba(k.get("Miesiąc pogrzebu")) or mp
        if not m_pogrz and m_zgonu and dp is not None:
            m_pogrz = m_zgonu % 12 + 1 if dz is not None and dp < dz else m_zgonu
        zgon = _data_lub_none(rok, m_zgonu, dz) if dz and m_zgonu else None
        pogrzeb = None
        if dp and m_pogrz:
            pogrzeb = _data_lub_none(rok + (1 if m_zgonu and m_pogrz < m_zgonu else 0), m_pogrz, dp)
        if (dz and m_zgonu and not zgon) or (dp and m_pogrz and not pogrzeb):
            raport.ostrz(nazwa_pliku, w, f"akt nr {nr}: niemożliwa data (zgon {dz}.{m_zgonu}, pogrzeb {dp}.{m_pogrz})")
        if zgon and pogrzeb and not 0 <= (dt.date.fromisoformat(pogrzeb) - dt.date.fromisoformat(zgon)).days <= 7:
            raport.ostrz(nazwa_pliku, w, f"akt nr {nr}: pogrzeb {pogrzeb} a zgon {zgon}")
        if (rok_z_daty(zgon) or rok_z_daty(pogrzeb) or rok) > min(args.max_rok, MAX_ROK_TYPU.get(typ, args.max_rok)):
            raport.odrzucone_rok[rok] += 1
            continue

        plec = {"M": "M", "K": "K"}.get((t("Płeć") or "").upper()[:1])
        religia = {"tak": "kat", "nie": "inna"}.get(zloz(t("Katolik") or ""))
        rodzaj = zloz(t("Rodzaj aktu") or "")

        # osoba zmarła: zamężne kobiety pod nazwiskiem panieńskim, dzieci bez nazwiska - po rodzicu
        imie, nazwisko_akt, rodowe = t("Zmarły – imię"), t("Zmarły – nazwisko"), t("Zmarły – nazwisko rodowe")
        oj_imie, oj_nazwisko = t("Ojciec – imię"), t("Ojciec – nazwisko")
        ma_imie, ma_rodowe = t("Matka – imię"), t("Matka – nazwisko rodowe")
        nazwisko = nazwisko_w_akcie = wg_rodzica = nazwisko_auto = None
        if plec == "K" and rodzaj in ("zona", "wdowa"):
            nazwisko = rodowe or oj_nazwisko
            nazwisko_auto = True if nazwisko and not rodowe else None
            if nazwisko_akt and zloz(nazwisko_akt) != zloz(nazwisko or ""):
                nazwisko_w_akcie = nazwisko_akt
        else:
            nazwisko = nazwisko_akt or oj_nazwisko or (ma_rodowe if "niesl" in rodzaj else None)
            if nazwisko and not nazwisko_akt and plec == "K" and nazwisko == oj_nazwisko:
                nazwisko = forma_zenska(nazwisko)
            if nazwisko and (not nazwisko_akt or (rodzaj.startswith("dziecko") or rodzaj.startswith("martwe"))
                             and not zloz(opis).startswith(zloz(f"{imie} {nazwisko}"))):
                wg_rodzica = True  # nazwisko nie stoi przy imieniu w akcie, tylko przy rodzicu
        if not nazwisko:
            raport.bez_nazwiska.append(f"{nazwa_pliku} w.{w}: {opis[:70]}")
        mz_imie, mz_nazwisko, mz_rodowe = t("Małżonek – imię"), t("Małżonek – nazwisko"), t("Małżonek – nazwisko rodowe")
        malzonek = " ".join(s for s in (mz_imie, mz_nazwisko) if s) or None
        if malzonek and mz_rodowe:
            malzonek += f" z domu {mz_rodowe}"

        # przyczyna: sama choroba; kapłan: samo imię i nazwisko (kolumna "Ksiądz");
        # z dopisków do "dodatkowych" idzie tylko to, co nie jest formułą o sakramentach ani o pochowaniu/zapisaniu
        przyczyna = t("Przyczyna zgonu")
        if przyczyna:  # kilka przyczyn bez przecinka ("Starość Zapalenie") -> "starość, zapalenie"
            przyczyna = ", ".join(c[0].lower() + c[1:] if not c[:2].isupper() else c
                                  for c in re.split(rf"\s+(?=[{CAP}])", przyczyna))
        reszta, kaplan_zdanie = rozdziel_kaplana(t("Adnotacje / dopiski"))
        kaplan = t("Ksiądz") or kaplan_zdanie
        if reszta:
            zdania = re.split(r"(?<=\.)\s+", reszta)
            reszta = " ".join(z for z in zdania if not re.match(r"(Zapisałem|Wpisałem|Pochowałem)\b", z)
                              and "akrament" not in z) or None
        inne = t("Inne informacje z aktu")
        if inne and zloz(inne.rstrip(".")) in zloz(opis):
            inne = None  # to samo, co już stoi w tekście aktu
        dodatkowe = poprawione(" ".join(s for s in (inne, reszta) if s) or None)
        wiek_tekst = _bez_ulamka_lat(t("Wiek"))

        rekord = {
            "_wiersz": w,
            "parafia": parafia,
            "typ": typ,
            "rok": rok,
            "nr": nr,
            "data_zgonu": zgon,
            "data_pogrzebu": pogrzeb,
            "dom": t("Numer domu"),
            "imie": imie,
            "nazwisko": nazwisko,
            "nazwisko_auto": nazwisko_auto,
            "nazwisko_wg_rodzica": wg_rodzica,
            "nazwisko_w_akcie": nazwisko_w_akcie,
            "imie_ojca": oj_imie,
            "imie_matki": ma_imie,
            "nazwisko_matki": ma_rodowe,
            "malzonek": malzonek,
            "plec": plec,
            "religia": religia,
            "wiek": wiek_pelne_lata(wiek_tekst),
            "wiek_lat": wiek_w_latach(wiek_tekst),
            "przyczyna": poprawione(przyczyna),
            "opis": poprawione(opis),
            "dodatkowe": dodatkowe,
            "kaplan": poprawione(kaplan),
        }
        rekordy.append({k: v for k, v in rekord.items() if v is not None})
    return rekordy


def wczytaj_plik(plik: Path, args, raport: Raport, znane: set[str] | None = None) -> list[dict]:
    znane = znane or set()
    parafia, typ, rok_folder = rozpoznaj_sciezke(plik)
    if not parafia or not typ:
        raport.pliki_pominiete.append(f"{plik}: nie da się ustalić parafii/rodzaju ze ścieżki (oczekiwane <Parafia>/<Zgony>/<rok>/plik.xlsx)")
        return []
    if typ not in TYPY_OBSLUGIWANE:
        raport.pliki_pominiete.append(f"{plik.name}: rodzaj '{typ}' nie jest jeszcze obsługiwany")
        return []

    wb = openpyxl.load_workbook(plik, data_only=True)
    if czy_weryfikacja(wb):
        return wczytaj_weryfikacje(wb, plik, parafia, typ, rok_folder, args, raport)
    ws = wb.worksheets[0]
    nazwa_pliku = f"{parafia}/{plik.parent.name}/{plik.name}"
    try:
        kol = wykryj_kolumny(ws)
    except BladUkladu as e:
        raport.pliki_pominiete.append(f"{nazwa_pliku}: {e}")
        return []

    m = re.match(r"\s*(\d{4})", str(ws["B1"].value or ""))
    rok_naglowka = int(m[1]) if m else None
    rok = rok_naglowka or rok_folder
    if rok_naglowka and rok_folder and rok_naglowka != rok_folder:
        raport.ostrz(nazwa_pliku, None, f"rok w nagłówku ({rok_naglowka}) różni się od folderu ({rok_folder}); użyto {rok_naglowka}")
    if not rok:
        raport.pliki_pominiete.append(f"{nazwa_pliku}: nie ustalono roku (brak w nagłówku B1 i w nazwie folderu)")
        return []

    def komorka(wiersz: int, nazwa: str):
        c = kol.get(nazwa)
        return ws.cell(wiersz, c).value if c else None

    rekordy: list[dict] = []
    poprzedni_pogrzeb: str | None = None
    for w in range(3, ws.max_row + 1):
        nr = liczba(komorka(w, "nr"))
        opis = tekst(komorka(w, "opis"))
        if nr is None and opis is None:
            continue
        if opis is None:
            raport.ostrz(nazwa_pliku, w, f"akt nr {nr} bez opisu – pominięty")
            continue

        zgon, p1 = parsuj_date(komorka(w, "data_zgonu"))
        pogrzeb, p2 = parsuj_date(komorka(w, "data_pogrzebu"))
        slub, p3 = parsuj_date(komorka(w, "data_slubu"))  # późniejsza adnotacja o ślubie (jeśli jest kolumna)
        for p in (p1, p2, p3):
            if p:
                raport.ostrz(nazwa_pliku, w, p)

        data_ur, wiek_tekst = None, None
        for c in kol["wiekowe"]:
            d, t = rozdziel_wiek(ws.cell(w, c).value)
            data_ur = data_ur or d
            wiek_tekst = wiek_tekst or t

        rok_zdarzenia = rok_z_daty(zgon) or rok_z_daty(pogrzeb) or rok
        if rok_zdarzenia > min(args.max_rok, MAX_ROK_TYPU.get(typ, args.max_rok)):
            raport.odrzucone_rok[rok_zdarzenia] += 1
            continue

        for etykieta, d in (("zgonu", zgon), ("pogrzebu", pogrzeb)):
            if d and len(d) == 10 and abs(int(d[:4]) - rok) > 1:
                raport.ostrz(nazwa_pliku, w, f"data {etykieta} {d} nie pasuje do roku księgi {rok}")
        if pogrzeb and len(pogrzeb) == 10:
            if poprzedni_pogrzeb and pogrzeb < poprzedni_pogrzeb:
                cofniecie = (dt.date.fromisoformat(poprzedni_pogrzeb) - dt.date.fromisoformat(pogrzeb)).days
                if cofniecie > 31:
                    raport.ostrz(nazwa_pliku, w, f"data pogrzebu {pogrzeb} cofa się o {cofniecie} dni względem poprzedniego wpisu (zamieniony dzień z miesiącem?)")
            poprzedni_pogrzeb = pogrzeb

        plec = None
        if zaznaczone(komorka(w, "plec_m")) and not zaznaczone(komorka(w, "plec_k")):
            plec = "M"
        elif zaznaczone(komorka(w, "plec_k")) and not zaznaczone(komorka(w, "plec_m")):
            plec = "K"
        elif zaznaczone(komorka(w, "plec_m")) and zaznaczone(komorka(w, "plec_k")):
            raport.ostrz(nazwa_pliku, w, "zaznaczona i płeć męska, i żeńska")
        religia = "kat" if zaznaczone(komorka(w, "rel_kat")) else ("inna" if zaznaczone(komorka(w, "rel_inna")) else None)

        # ---- teksty: bez łaciny w nawiasach, miejscowości w miejscowniku (ur. w Przytkowicach)
        opis_surowy = opis
        opis = poprawione(opis_surowy)
        przyczyna = tekst(komorka(w, "przyczyna"))
        reszta_przyczyny = None
        if przyczyna and " | " in przyczyna:  # starszy układ: "przyczyna | zaopatrzenie, kto pochował"
            przyczyna, reszta_przyczyny = (s.strip() for s in przyczyna.split(" | ", 1))
        przyczyna, reszta_przyczyny = poprawione(przyczyna), poprawione(reszta_przyczyny)
        uwagi_jawne = None
        if not args.bez_uwag:
            czesci = [tekst(ws.cell(w, c).value) for c in kol["uwagi"]]
            uwagi_jawne = poprawione(" | ".join(t for t in czesci if t) or None)
        kaplan_jawny = poprawione(tekst(komorka(w, "kaplan")))

        # ---- osoba zmarła
        z_opisu = rozbierz_opis(opis_surowy)
        imie, nazwisko = z_opisu.get("imie"), z_opisu.get("nazwisko")
        wg_rodzica = z_opisu.get("nazwisko_wg_rodzica")
        if wg_rodzica and nazwisko:  # nazwisko z opisu rodziców bywa w dopełniaczu (Żaka -> Żak)
            nazwisko = nazwisko_mianownik(nazwisko, znane)
        opisowe = nazwisko  # nazwisko w opisie aktu (u zamężnych kobiet – po mężu)
        nazwisko_w_akcie = nazwisko_auto = None
        jawna = tekst(komorka(w, "osoba"))  # kolumna "Imię nazwisko" w nowym układzie
        matka_j, ojciec_j, malzonek_j = (tekst(komorka(w, k)) for k in ("matka", "ojciec", "malzonek"))
        rodz = rodzina_z_opisu(opis_surowy, opisowe, plec, znane)
        zamezna = plec == "K" and kobieta_zamezna(opis_surowy)
        if jawna:
            imie, nazwisko = rozdziel_imie_nazwisko(jawna)
            if len(jawna.split()) > 3:
                raport.ostrz(nazwa_pliku, w, f"w kolumnie 'Imię nazwisko' jest więcej niż 3 wyrazy: '{jawna}'")
            if opisowe and not wg_rodzica and (not nazwisko or zloz(opisowe) != zloz(nazwisko)):
                nazwisko_w_akcie = opisowe  # np. nazwisko po mężu użyte w opisie aktu
            wg_rodzica = None
        elif zamezna:
            # kobiety pod nazwiskiem panieńskim: wprost z opisu ("z domu X") albo = nazwisko ojca;
            # gdy opis nie podaje ani jednego, ani drugiego – bez nazwiska
            nazwisko = nazwisko_panienskie(opis_surowy)
            if not nazwisko:
                nazwisko = rodz.get("nazwisko_ojca")
                nazwisko_auto = True if nazwisko else None
            wg_rodzica = None
            if opisowe and (not nazwisko or zloz(opisowe) != zloz(nazwisko)):
                nazwisko_w_akcie = opisowe
        if not nazwisko and not zamezna:
            raport.bez_nazwiska.append(f"{nazwa_pliku} w.{w}: {opis[:70]}")

        # ---- rodzice i małżonek: z kolumn Excela, a gdy wiersz nie jest jeszcze przerobiony – wstępnie z opisu
        imie_ojca = imie_matki = nazwisko_matki = None
        malzonek = malzonek_j
        rodzice_auto = None
        if jawna or matka_j or ojciec_j or malzonek_j:
            if ojciec_j:
                imie_ojca = rozdziel_imie_nazwisko(ojciec_j)[0]
            if matka_j:
                imie_matki, nazwisko_matki = rozdziel_imie_nazwisko(matka_j)
            # komórka zaznaczona na żółto przez podpowiedz_excel.py = jeszcze niesprawdzona (strona pokaże ją kursywą)
            zolte = {n for n in ("osoba", "matka", "ojciec", "malzonek") if n in kol and zolta(ws.cell(w, kol[n]))}
            if zolte & {"matka", "ojciec", "malzonek"}:
                rodzice_auto = True
            if "osoba" in zolte and nazwisko:
                nazwisko_auto = True
        else:
            imie_ojca, imie_matki, nazwisko_matki = rodz.get("imie_ojca"), rodz.get("imie_matki"), rodz.get("nazwisko_matki")
            malzonek = rodz.get("malzonek")
            rodzice_auto = True if (imie_ojca or imie_matki or malzonek) else None

        # ---- "Dodatkowe informacje" (bez kapłana) i "Kapłan"
        if jawna or kaplan_jawny:
            reszta, kaplan = uwagi_jawne, kaplan_jawny  # wiersz przerobiony ręcznie: bierzemy to, co wpisano w Excelu
        else:
            # stary układ: reszta opisu od "ur. w ...", reszta z kolumny przyczyny -> rozdzielone na uwagi i kapłana;
            # notatki tłumacza (kolumna bez nagłówka) dopisujemy bez dzielenia
            reszta, kaplan = rozdziel_kaplana(" ".join(t for t in (dodatkowe_z_opisu(opis), reszta_przyczyny) if t) or None)
            if uwagi_jawne:
                reszta = f"{reszta} {uwagi_jawne}" if reszta else uwagi_jawne
        dodatkowe = reszta

        skan = tekst(komorka(w, "skan"))
        if skan and not skan.lower().startswith("http"):
            skan = normalizuj_skan(skan)

        rekord = {
            "_wiersz": w,
            "parafia": parafia,
            "typ": typ,
            "rok": rok,
            "nr": nr,
            "data_zgonu": zgon,
            "data_pogrzebu": pogrzeb,
            "data_slubu": slub,
            "dom": tekst(komorka(w, "dom")),
            "imie": imie,
            "nazwisko": nazwisko,
            "nazwisko_auto": nazwisko_auto,
            "nazwisko_wg_rodzica": wg_rodzica,
            "nazwisko_w_akcie": nazwisko_w_akcie,
            "imie_ojca": imie_ojca,
            "imie_matki": imie_matki,
            "nazwisko_matki": nazwisko_matki,
            "rodzice_auto": rodzice_auto,
            "matka": matka_j,
            "ojciec": ojciec_j,
            "malzonek": malzonek,
            "plec": plec,
            "religia": religia,
            "data_ur": data_ur,
            "wiek": wiek_pelne_lata(wiek_tekst),
            "wiek_lat": wiek_w_latach(wiek_tekst),
            "przyczyna": przyczyna,
            "opis": opis,
            "dodatkowe": dodatkowe,
            "kaplan": kaplan,
            "skan": skan,
            # dane pomocnicze dla podpowiedz_excel.py (nie trafiają do JSON)
            "_surowe": {
                "opis_surowy": opis_surowy, "zgon": komorka(w, "data_zgonu"), "pogrzeb": komorka(w, "data_pogrzebu"),
                "wiek_tekst": wiek_tekst, "jawna": jawna, "rodz": rodz, "zamezna": zamezna, "opisowe": opisowe,
                "uwagi_jawne": uwagi_jawne, "kaplan_jawny": kaplan_jawny, "skan": skan, "przyczyna": przyczyna,
            },
        }
        rekordy.append({k: v for k, v in rekord.items() if v is not None})
    return rekordy


# ---------------------------------------------------------------- zapis

def zbierz_pliki(sciezki: list[Path]) -> list[Path]:
    pliki: list[Path] = []
    for s in sciezki:
        if s.is_file():
            pliki.append(s)
        elif s.is_dir():
            pliki.extend(p for p in s.rglob("*.xlsx") if "_nowy" not in p.stem)  # pomijamy pliki z podpowiedz_excel.py
        else:
            sys.exit(f"Nie znaleziono: {s}")
    return sorted({p.resolve() for p in pliki if not p.name.startswith("~$")})


def zapisz_json(sciezka: Path, naglowek: dict, rekordy: list[dict]) -> None:
    glowa = json.dumps(naglowek, ensure_ascii=False, separators=(",", ":"))[:-1]  # bez zamykającego }
    wiersze = ",\n".join(json.dumps(r, ensure_ascii=False, separators=(",", ":")) for r in rekordy)
    sciezka.write_text(f'{glowa},"rekordy":[\n{wiersze}\n]}}\n', encoding="utf-8")


def wczytaj_liste_zrodel() -> list[Path]:
    if not PLIK_LISTY_ZRODEL.exists():
        sys.exit(
            "Nie podano źródeł. Użyj --zrodlo <plik lub folder> albo utwórz plik "
            f"{PLIK_LISTY_ZRODEL.relative_to(KORZEN)} (po jednej ścieżce w wierszu; wzór: zrodla.przyklad.txt)."
        )
    wiersze = PLIK_LISTY_ZRODEL.read_text(encoding="utf-8").splitlines()
    return [Path(w.strip().strip('"')) for w in wiersze if w.strip() and not w.strip().startswith("#")]


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--zrodlo", action="append", type=Path, help="plik .xlsx lub folder (można podać kilka razy)")
    ap.add_argument("--max-rok", type=int, default=MAX_ROK_DOMYSLNIE, help=f"nie publikuj aktów późniejszych niż ten rok (domyślnie {MAX_ROK_DOMYSLNIE})")
    ap.add_argument("--bez-uwag", action="store_true", help="nie publikuj kolumny 'Uwagi' z Excela")
    ap.add_argument("--sucho", action="store_true", help="tylko sprawdź i pokaż raport, nic nie zapisuj")
    ap.add_argument("--szczegoly", action="store_true", help="pokaż wszystkie ostrzeżenia i rekordy bez rozpoznanego nazwiska")
    ap.add_argument("--podglad", type=int, metavar="N", help="wypisz N pierwszych rekordów jako JSON (do sprawdzenia)")
    args = ap.parse_args()

    pliki = zbierz_pliki(args.zrodlo or wczytaj_liste_zrodel())
    if not pliki:
        sys.exit("Nie znaleziono żadnych plików .xlsx.")

    raport = Raport()
    znane = zbierz_znane(pliki)
    zbiory: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for p in pliki:
        for r in wczytaj_plik(p, args, raport, znane):
            zbiory[(r["parafia"], r["typ"])].append(r)

    # identyfikatory (unikalne nawet gdy numer aktu się powtórzy między plikami)
    uzyte: Counter = Counter()
    for (parafia, typ), rekordy in zbiory.items():
        rekordy.sort(key=lambda r: (r["rok"], r.get("nr") is None, r.get("nr") or 0, r.get("data_pogrzebu") or ""))
        for r in rekordy:
            if r.get("nr") is None:
                raport.ostrz(f"{parafia}/{r['rok']}", r["_wiersz"], "akt bez numeru – id z numeru wiersza")
                baza = f"{slug(parafia)}-{typ}-{r['rok']}-w{r['_wiersz']}"
            else:
                baza = f"{slug(parafia)}-{typ}-{r['rok']}-{r['nr']}"
            uzyte[baza] += 1
            r["id"] = baza if uzyte[baza] == 1 else f"{baza}-{uzyte[baza]}"
            if uzyte[baza] > 1:
                raport.ostrz(f"{parafia}/{r['rok']}", None, f"powtórzony numer aktu {r.get('nr')} – nadano id {r['id']}")
        # id na początek rekordu, bez pól roboczych
        for i, r in enumerate(rekordy):
            rekordy[i] = {"id": r["id"], **{k: v for k, v in r.items() if k != "id" and not k.startswith("_")}}

    # ----- raport
    razem = sum(len(v) for v in zbiory.values())
    print(f"Pliki Excel: {len(pliki)}  |  rekordy: {razem}  |  limit roku: {args.max_rok}")
    for (parafia, typ), rekordy in sorted(zbiory.items()):
        lata = Counter(r["rok"] for r in rekordy)
        print(f"  {parafia} / {typ}: {len(rekordy)} rek.  lata: " + ", ".join(f"{k} ({v})" for k, v in sorted(lata.items())))
    if raport.odrzucone_rok:
        print("Odrzucone (późniejsze niż limit): " + ", ".join(f"{k}: {v}" for k, v in sorted(raport.odrzucone_rok.items())))
    for p in raport.pliki_pominiete:
        print("POMINIĘTO", p)
    bez = len(raport.bez_nazwiska)
    print(f"Bez rozpoznanego nazwiska: {bez} z {razem}")
    wszystkie = [r for v in zbiory.values() for r in v]
    z_excela = sum(1 for r in wszystkie if (r.get("matka") or r.get("ojciec") or r.get("imie_ojca") or r.get("imie_matki")
                                             or r.get("malzonek")) and not r.get("rodzice_auto"))
    z_opisu_n = sum(1 for r in wszystkie if r.get("rodzice_auto"))
    print(f"Rodzice: z kolumn Excela {z_excela}, wstępnie (z opisu lub żółte do sprawdzenia) {z_opisu_n}, brak {razem - z_excela - z_opisu_n}"
          f"  |  imię ojca: {sum(1 for r in wszystkie if r.get('imie_ojca'))}, imię matki: {sum(1 for r in wszystkie if r.get('imie_matki'))},"
          f" nazwisko matki: {sum(1 for r in wszystkie if r.get('nazwisko_matki'))}")
    print(f"Ostrzeżenia: {len(raport.ostrzezenia)}")
    limit = None if args.szczegoly else 15
    for o in raport.ostrzezenia[:limit]:
        print("  !", o)
    if args.szczegoly:
        for b in raport.bez_nazwiska:
            print("  ?", b)
    elif bez:
        for b in raport.bez_nazwiska[:5]:
            print("  ?", b)
        print("  (--szczegoly pokaże wszystkie)")

    if args.podglad:
        for r in [r for v in zbiory.values() for r in v][: args.podglad]:
            print(json.dumps(r, ensure_ascii=False, indent=1))

    if args.sucho:
        print("Tryb --sucho: nic nie zapisano.")
        return
    if not razem:
        sys.exit("Brak rekordów do zapisania – nic nie zapisano.")

    KATALOG_WYJSCIA.mkdir(parents=True, exist_ok=True)
    for stary in KATALOG_WYJSCIA.glob("*.json"):
        stary.unlink()
    manifest_zbiory = []
    for (parafia, typ), rekordy in sorted(zbiory.items()):
        nazwa_pliku = f"{slug(parafia)}-{typ}.json"
        zapisz_json(KATALOG_WYJSCIA / nazwa_pliku, {"parafia": parafia, "typ": typ}, rekordy)
        manifest_zbiory.append({
            "plik": nazwa_pliku,
            "parafia": parafia,
            "parafia_id": slug(parafia),
            "typ": typ,
            "lata": sorted({r["rok"] for r in rekordy}),
            "rekordow": len(rekordy),
        })
    zrodla = {}
    if PLIK_ZRODEL_SKANOW.exists():
        zrodla = json.loads(PLIK_ZRODEL_SKANOW.read_text(encoding="utf-8"))
    manifest = {
        "wygenerowano": dt.date.today().isoformat(),
        "max_rok": args.max_rok,
        "max_rok_typ": MAX_ROK_TYPU,
        "licencja": LICENCJA,
        "zbiory": manifest_zbiory,
        "zrodla": zrodla,
    }
    (KATALOG_WYJSCIA / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Zapisano {len(manifest_zbiory)} zbiór(y) + manifest.json w {KATALOG_WYJSCIA.relative_to(KORZEN)}")


if __name__ == "__main__":
    main()
